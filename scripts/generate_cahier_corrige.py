from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "cahier-des-charges-jobconnect-corrige.docx"
OUT_MD = ROOT / "cahier-des-charges-jobconnect-corrige.md"
OUT_PDF = ROOT / "cahier-des-charges-jobconnect-corrige.pdf"


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(26, 32, 44)
MUTED = RGBColor(91, 103, 112)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F4F6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.color.rgb = BLUE if level <= 2 else DARK
    if level == 1:
        run.font.size = Pt(16)
    elif level == 2:
        run.font.size = Pt(13)
    else:
        run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(hdr[i], "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], "F8FAFC")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build_markdown():
    return """# Cahier des charges - JobConnect

## Informations generales

**Nom du projet :** JobConnect

**Sujet :** Plateforme web de recrutement pour candidats et entreprises

**Candidat :** Saad Berra

**Etablissement :** SupTechnology

**Filiere :** Developpement Web et Mobile

**Annee universitaire :** 2025-2026

**Encadrant :** M. Ahmed Zellou

**Objectif du document :** definir le besoin, le perimetre, les acteurs, les fonctionnalites attendues, les contraintes et les criteres de validation du projet JobConnect.

## 1. Introduction

JobConnect est une plateforme web de recrutement destinee a faciliter la relation entre les candidats a la recherche d'opportunites professionnelles et les entreprises souhaitant publier des offres et suivre les candidatures.

Le document presente le cahier des charges fonctionnel du projet. Il decrit les besoins a satisfaire, les roles des utilisateurs, les regles de gestion principales, les contraintes non fonctionnelles et les criteres permettant de valider la solution.

## 2. Contexte du projet

Le recrutement en ligne est devenu un moyen essentiel pour diffuser les offres d'emploi, recevoir les candidatures et gerer les echanges entre recruteurs et candidats. Cependant, les plateformes classiques peuvent parfois manquer de controle sur la fiabilite des comptes, la pertinence des candidatures et la securite des donnees personnelles.

JobConnect propose une solution organisee autour de trois espaces : candidat, entreprise et administrateur. Elle permet de centraliser les profils, les offres, les candidatures, les CV, les suggestions et la messagerie dans un environnement controle.

## 3. Etude de l'existant

Plusieurs plateformes de recrutement existent deja, telles que LinkedIn, Indeed, ReKrute, Emploi.ma et ANAPEC. Elles proposent des fonctionnalites importantes comme la publication d'offres, le depot de CV, la recherche de profils et les alertes.

Ces solutions restent souvent generalistes ou orientees vers un usage massif. Dans le cadre de ce projet, JobConnect se distingue par une approche encadree : validation des entreprises, controle des acces, protection des CV, messagerie limitee aux relations autorisees et suivi administratif des signalements.

| Critere | Plateformes existantes | JobConnect |
| --- | --- | --- |
| Publication d'offres | Disponible | Disponible pour les entreprises validees |
| Espace candidat | Disponible | Profil complet, CV, competences et suivi |
| Espace entreprise | Disponible | Publication, candidatures et suggestions |
| Validation entreprise | Variable | Validation obligatoire par administrateur |
| Matching | Souvent partiel | Score selon profil, competences et offre |
| Moderation | Variable | Blocage, signalements et suivi admin |

## 4. Problematique

Le projet repond aux questions suivantes :

- comment aider un candidat a trouver des offres adaptees a son profil ?
- comment aider une entreprise a identifier rapidement les candidatures pertinentes ?
- comment proteger les CV et les donnees personnelles ?
- comment limiter les abus dans les echanges entre candidats et recruteurs ?
- comment permettre a l'administrateur de controler les comptes, les offres et les signalements ?

## 5. Objectifs du projet

### 5.1 Objectif principal

Concevoir et realiser une application web de recrutement permettant une mise en relation securisee, organisee et pertinente entre candidats et entreprises.

### 5.2 Objectifs specifiques

- permettre l'inscription et la connexion des utilisateurs ;
- gerer trois espaces distincts : candidat, entreprise et administrateur ;
- permettre aux entreprises validees de publier et gerer leurs offres ;
- permettre aux candidats de completer leur profil, gerer leur CV et postuler ;
- proposer un systeme de matching entre offres et candidats ;
- assurer le suivi des candidatures et des notifications ;
- encadrer la messagerie pour eviter les abus ;
- fournir un espace d'administration pour la moderation ;
- garantir la securite des donnees sensibles.

## 6. Perimetre du projet

### 6.1 Fonctionnalites incluses

- gestion des comptes utilisateurs ;
- verification des comptes et controle des roles ;
- gestion des profils candidats ;
- gestion des profils entreprises ;
- publication et gestion des offres ;
- depot, consultation et generation de CV ;
- gestion des candidatures ;
- matching et suggestions ;
- notifications ;
- messagerie controlee ;
- signalements ;
- tableaux de bord ;
- moderation administrative.

### 6.2 Fonctionnalites exclues

- paiement en ligne ;
- application mobile native ;
- visioconference ;
- signature electronique de contrat ;
- integration directe avec des reseaux sociaux externes ;
- systeme d'intelligence artificielle externe.

## 7. Acteurs du systeme

### 7.1 Visiteur

Le visiteur peut consulter les offres publiques, filtrer les offres, voir le detail d'une offre, creer un compte et se connecter.

### 7.2 Candidat

Le candidat peut creer un compte, completer son profil, ajouter ses competences, deposer ou generer un CV, consulter les offres, postuler, suivre ses candidatures, recevoir des notifications et communiquer lorsque les regles de messagerie l'autorisent.

### 7.3 Entreprise

L'entreprise peut creer un compte, attendre la validation administrative, completer son profil, publier des offres apres validation, gerer ses offres, consulter les candidatures, changer les statuts, consulter les suggestions et communiquer avec les candidats autorises.

### 7.4 Administrateur

L'administrateur gere les utilisateurs, valide ou refuse les entreprises, bloque les comptes si necessaire, gere les categories et competences, consulte les offres, traite les signalements et supervise l'activite globale.

## 8. Besoins fonctionnels

### 8.1 Gestion des utilisateurs

| Besoin | Description | Priorite |
| --- | --- | --- |
| Inscription | Creation d'un compte candidat ou entreprise | Haute |
| Connexion | Acces securise a l'espace utilisateur | Haute |
| Verification | Confirmation de l'adresse email | Haute |
| Roles | Separation des droits candidat, entreprise et admin | Haute |
| Blocage | Suspension possible d'un compte par l'admin | Haute |

### 8.2 Espace candidat

| Besoin | Description | Priorite |
| --- | --- | --- |
| Profil | Informations personnelles et professionnelles | Haute |
| Competences | Ajout des competences avec niveau | Haute |
| CV | Depot et generation d'un CV | Haute |
| Candidatures | Postuler et suivre les statuts | Haute |
| Suggestions | Recevoir des offres adaptees | Moyenne |
| Messagerie | Echanger selon les regles autorisees | Moyenne |

### 8.3 Espace entreprise

| Besoin | Description | Priorite |
| --- | --- | --- |
| Profil entreprise | Informations de presentation de l'entreprise | Haute |
| Offres | Creation, modification et suppression d'offres | Haute |
| Candidatures | Consultation et traitement des candidatures | Haute |
| CV candidats | Consultation securisee des CV recus | Haute |
| Suggestions | Liste de candidats pertinents par offre | Moyenne |
| Signalement | Declaration d'un abus a l'administrateur | Moyenne |

### 8.4 Espace administrateur

| Besoin | Description | Priorite |
| --- | --- | --- |
| Dashboard | Statistiques globales de la plateforme | Haute |
| Utilisateurs | Consultation, blocage et deblocage | Haute |
| Entreprises | Validation ou refus des comptes entreprises | Haute |
| Referentiels | Gestion des categories et competences | Moyenne |
| Signalements | Traitement des signalements | Moyenne |
| Moderation | Supervision des contenus et conversations | Moyenne |

## 9. Regles de gestion principales

- une entreprise doit etre validee avant de publier des offres ;
- un candidat ne peut postuler qu'une seule fois a la meme offre ;
- une candidature possede un statut de suivi : en attente, en cours, acceptee ou refusee ;
- les CV doivent etre proteges et accessibles uniquement aux utilisateurs autorises ;
- une entreprise peut contacter un candidat dans le cadre d'une candidature ;
- un candidat peut repondre ou contacter l'entreprise uniquement lorsque la relation est autorisee ;
- l'administrateur peut bloquer un utilisateur et traiter les signalements ;
- les offres d'une entreprise refusee, bloquee ou non validee ne doivent pas etre visibles publiquement.

## 10. Matching et suggestions

Le matching permet d'estimer la compatibilite entre une offre et un candidat. Il repose sur les informations du profil, les competences, la localisation, le niveau d'etude et les informations disponibles dans le CV lorsque celles-ci sont exploitables.

L'objectif n'est pas de remplacer le recruteur, mais de l'aider a prioriser les profils et d'aider le candidat a identifier les offres les plus pertinentes.

## 11. Besoins non fonctionnels

### 11.1 Securite

- protection des mots de passe ;
- controle des acces selon les roles ;
- protection contre les actions non autorisees ;
- stockage prive des CV ;
- validation des donnees saisies ;
- limitation des tentatives sensibles ;
- moderation des messages et signalements.

### 11.2 Ergonomie

- interface claire pour chaque type d'utilisateur ;
- navigation simple entre les principales fonctionnalites ;
- messages d'erreur et de confirmation comprehensibles ;
- tableaux de bord lisibles ;
- formulaires organises et faciles a remplir.

### 11.3 Performance et fiabilite

- listes paginees lorsque les donnees sont nombreuses ;
- temps de reponse acceptable pour les operations courantes ;
- conservation correcte des donnees ;
- disponibilite de donnees de demonstration pour tester les parcours.

### 11.4 Maintenabilite

- organisation claire du code ;
- separation des responsabilites ;
- facilite d'evolution des fonctionnalites ;
- validation du comportement par des tests.

## 12. Contraintes du projet

- le projet doit etre une application web utilisable en local pour la demonstration ;
- la base de donnees doit stocker les comptes, profils, offres, candidatures, messages et signalements ;
- les documents sensibles, notamment les CV, ne doivent pas etre exposes publiquement ;
- la solution doit etre compatible avec le contexte de formation en developpement web et mobile ;
- les fonctionnalites doivent rester demonstrables pendant la soutenance.

## 13. Criteres d'acceptation

Le projet sera considere valide si :

- un visiteur peut consulter les offres publiques ;
- un candidat peut creer un compte, completer son profil, ajouter un CV et postuler ;
- une entreprise validee peut publier une offre et gerer les candidatures ;
- une entreprise non validee ne peut pas publier d'offre ;
- le systeme propose des suggestions ou scores de matching ;
- les CV ne sont accessibles qu'aux utilisateurs autorises ;
- la messagerie respecte les regles de controle ;
- l'administrateur peut gerer les utilisateurs, entreprises, referentiels et signalements ;
- les principaux parcours peuvent etre presentes sans erreur lors de la demonstration.

## 14. Planning previsionnel

| Phase | Contenu | Etat |
| --- | --- | --- |
| Phase 1 | Analyse du besoin et conception | Terminee |
| Phase 2 | Mise en place du projet et authentification | Terminee |
| Phase 3 | Espace candidat et gestion du CV | Terminee |
| Phase 4 | Espace entreprise et gestion des offres | Terminee |
| Phase 5 | Espace administrateur et moderation | Terminee |
| Phase 6 | Matching, suggestions et notifications | Terminee |
| Phase 7 | Tests, corrections et documentation | En cours |
| Phase 8 | Preparation de la soutenance et finalisation | En cours |

## 15. Evolutions possibles

- amelioration du score de matching ;
- recherche avancee multicritere ;
- calendrier d'entretien ;
- historique detaille des actions administratives ;
- export PDF avance du CV ;
- application mobile ;
- statistiques plus detaillees pour les entreprises.

## 16. Conclusion

JobConnect vise a fournir une plateforme de recrutement complete, securisee et adaptee aux besoins des candidats, des entreprises et de l'administrateur. Le projet met l'accent sur la qualite des parcours utilisateurs, la protection des donnees, la moderation et l'aide a la decision grace au matching.

Ce cahier des charges servira de reference pour verifier que la solution realisee repond bien aux besoins fonctionnels et non fonctionnels definis.
"""


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "JobConnect - Cahier des charges corrige"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "Projet de Fin d'Etudes - 2025-2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(80)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("JobConnect")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = BLUE

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Cahier des charges fonctionnel")
    r.font.size = Pt(17)
    r.font.color.rgb = DARK
    r.bold = True

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = desc.add_run("Plateforme web de recrutement pour candidats et entreprises")
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = MUTED

    doc.add_paragraph()
    meta_rows = [
        ("Candidat", "Saad Berra"),
        ("Etablissement", "SupTechnology"),
        ("Filiere", "Developpement Web et Mobile"),
        ("Annee universitaire", "2025-2026"),
        ("Encadrant", "M. Ahmed Zellou"),
        ("Type de projet", "Application web de recrutement"),
        ("Objectif du document", "Definir le besoin, le perimetre, les acteurs, les fonctionnalites attendues, les contraintes et les criteres de validation."),
    ]
    add_table(doc, ["Information", "Valeur"], meta_rows, [4.5, 11.0])

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    nr = note.add_run("Version corrigee selon les remarques de l'encadrant : contenu recentre sur le cahier des charges, details techniques reduits.")
    nr.font.size = Pt(10)
    nr.font.color.rgb = MUTED

    doc.add_page_break()

    add_heading(doc, "1. Introduction")
    add_para(doc, "JobConnect est une plateforme web de recrutement destinee a faciliter la relation entre les candidats a la recherche d'opportunites professionnelles et les entreprises souhaitant publier des offres et suivre les candidatures.")
    add_para(doc, "Ce document presente le cahier des charges fonctionnel du projet. Il decrit les besoins a satisfaire, les roles des utilisateurs, les regles de gestion principales, les contraintes non fonctionnelles et les criteres permettant de valider la solution.")

    add_heading(doc, "2. Contexte du projet")
    add_para(doc, "Le recrutement en ligne est devenu un moyen essentiel pour diffuser les offres d'emploi, recevoir les candidatures et gerer les echanges entre recruteurs et candidats. Cependant, les plateformes classiques peuvent parfois manquer de controle sur la fiabilite des comptes, la pertinence des candidatures et la securite des donnees personnelles.")
    add_para(doc, "JobConnect propose une solution organisee autour de trois espaces : candidat, entreprise et administrateur. Elle permet de centraliser les profils, les offres, les candidatures, les CV, les suggestions et la messagerie dans un environnement controle.")

    add_heading(doc, "3. Etude de l'existant")
    add_para(doc, "Plusieurs plateformes de recrutement existent deja, telles que LinkedIn, Indeed, ReKrute, Emploi.ma et ANAPEC. Elles proposent des fonctionnalites importantes comme la publication d'offres, le depot de CV, la recherche de profils et les alertes.")
    add_para(doc, "Ces solutions restent souvent generalistes ou orientees vers un usage massif. JobConnect se distingue par une approche encadree : validation des entreprises, controle des acces, protection des CV, messagerie limitee aux relations autorisees et suivi administratif des signalements.")
    add_table(
        doc,
        ["Critere", "Plateformes existantes", "JobConnect"],
        [
            ("Publication d'offres", "Disponible", "Disponible pour les entreprises validees"),
            ("Espace candidat", "Disponible", "Profil complet, CV, competences et suivi"),
            ("Espace entreprise", "Disponible", "Publication, candidatures et suggestions"),
            ("Validation entreprise", "Variable", "Validation obligatoire par administrateur"),
            ("Matching", "Souvent partiel", "Score selon profil, competences et offre"),
            ("Moderation", "Variable", "Blocage, signalements et suivi admin"),
        ],
        [4.2, 5.6, 5.6],
    )

    add_heading(doc, "4. Problematique")
    add_para(doc, "Le projet repond aux questions suivantes :")
    add_bullets(doc, [
        "comment aider un candidat a trouver des offres adaptees a son profil ?",
        "comment aider une entreprise a identifier rapidement les candidatures pertinentes ?",
        "comment proteger les CV et les donnees personnelles ?",
        "comment limiter les abus dans les echanges entre candidats et recruteurs ?",
        "comment permettre a l'administrateur de controler les comptes, les offres et les signalements ?",
    ])

    add_heading(doc, "5. Objectifs du projet")
    add_heading(doc, "5.1 Objectif principal", 2)
    add_para(doc, "Concevoir et realiser une application web de recrutement permettant une mise en relation securisee, organisee et pertinente entre candidats et entreprises.")
    add_heading(doc, "5.2 Objectifs specifiques", 2)
    add_bullets(doc, [
        "permettre l'inscription et la connexion des utilisateurs ;",
        "gerer trois espaces distincts : candidat, entreprise et administrateur ;",
        "permettre aux entreprises validees de publier et gerer leurs offres ;",
        "permettre aux candidats de completer leur profil, gerer leur CV et postuler ;",
        "proposer un systeme de matching entre offres et candidats ;",
        "assurer le suivi des candidatures et des notifications ;",
        "encadrer la messagerie pour eviter les abus ;",
        "fournir un espace d'administration pour la moderation ;",
        "garantir la securite des donnees sensibles.",
    ])

    add_heading(doc, "6. Perimetre du projet")
    add_heading(doc, "6.1 Fonctionnalites incluses", 2)
    add_bullets(doc, [
        "gestion des comptes utilisateurs ;",
        "verification des comptes et controle des roles ;",
        "gestion des profils candidats et entreprises ;",
        "publication et gestion des offres ;",
        "depot, consultation et generation de CV ;",
        "gestion des candidatures, matching, suggestions et notifications ;",
        "messagerie controlee, signalements, tableaux de bord et moderation administrative.",
    ])
    add_heading(doc, "6.2 Fonctionnalites exclues", 2)
    add_bullets(doc, [
        "paiement en ligne ;",
        "application mobile native ;",
        "visioconference ;",
        "signature electronique de contrat ;",
        "integration directe avec des reseaux sociaux externes ;",
        "systeme d'intelligence artificielle externe.",
    ])

    add_heading(doc, "7. Acteurs du systeme")
    add_table(
        doc,
        ["Acteur", "Role principal"],
        [
            ("Visiteur", "Consulter les offres publiques, filtrer les offres, voir le detail d'une offre, creer un compte et se connecter."),
            ("Candidat", "Completer son profil, ajouter ses competences, gerer son CV, postuler, suivre ses candidatures et communiquer lorsque les regles l'autorisent."),
            ("Entreprise", "Completer son profil, publier des offres apres validation, consulter les candidatures, changer les statuts et contacter les candidats autorises."),
            ("Administrateur", "Gerer les utilisateurs, valider ou refuser les entreprises, traiter les signalements et superviser l'activite globale."),
        ],
        [3.8, 11.5],
    )

    add_heading(doc, "8. Besoins fonctionnels")
    add_heading(doc, "8.1 Gestion des utilisateurs", 2)
    add_table(doc, ["Besoin", "Description", "Priorite"], [
        ("Inscription", "Creation d'un compte candidat ou entreprise", "Haute"),
        ("Connexion", "Acces securise a l'espace utilisateur", "Haute"),
        ("Verification", "Confirmation de l'adresse email", "Haute"),
        ("Roles", "Separation des droits candidat, entreprise et admin", "Haute"),
        ("Blocage", "Suspension possible d'un compte par l'admin", "Haute"),
    ], [4.0, 8.5, 2.5])
    add_heading(doc, "8.2 Espace candidat", 2)
    add_table(doc, ["Besoin", "Description", "Priorite"], [
        ("Profil", "Informations personnelles et professionnelles", "Haute"),
        ("Competences", "Ajout des competences avec niveau", "Haute"),
        ("CV", "Depot et generation d'un CV", "Haute"),
        ("Candidatures", "Postuler et suivre les statuts", "Haute"),
        ("Suggestions", "Recevoir des offres adaptees", "Moyenne"),
        ("Messagerie", "Echanger selon les regles autorisees", "Moyenne"),
    ], [4.0, 8.5, 2.5])
    add_heading(doc, "8.3 Espace entreprise", 2)
    add_table(doc, ["Besoin", "Description", "Priorite"], [
        ("Profil entreprise", "Informations de presentation de l'entreprise", "Haute"),
        ("Offres", "Creation, modification et suppression d'offres", "Haute"),
        ("Candidatures", "Consultation et traitement des candidatures", "Haute"),
        ("CV candidats", "Consultation securisee des CV recus", "Haute"),
        ("Suggestions", "Liste de candidats pertinents par offre", "Moyenne"),
        ("Signalement", "Declaration d'un abus a l'administrateur", "Moyenne"),
    ], [4.0, 8.5, 2.5])
    add_heading(doc, "8.4 Espace administrateur", 2)
    add_table(doc, ["Besoin", "Description", "Priorite"], [
        ("Dashboard", "Statistiques globales de la plateforme", "Haute"),
        ("Utilisateurs", "Consultation, blocage et deblocage", "Haute"),
        ("Entreprises", "Validation ou refus des comptes entreprises", "Haute"),
        ("Referentiels", "Gestion des categories et competences", "Moyenne"),
        ("Signalements", "Traitement des signalements", "Moyenne"),
        ("Moderation", "Supervision des contenus et conversations", "Moyenne"),
    ], [4.0, 8.5, 2.5])

    add_heading(doc, "9. Regles de gestion principales")
    add_bullets(doc, [
        "une entreprise doit etre validee avant de publier des offres ;",
        "un candidat ne peut postuler qu'une seule fois a la meme offre ;",
        "une candidature possede un statut de suivi : en attente, en cours, acceptee ou refusee ;",
        "les CV doivent etre proteges et accessibles uniquement aux utilisateurs autorises ;",
        "une entreprise peut contacter un candidat dans le cadre d'une candidature ;",
        "un candidat peut repondre ou contacter l'entreprise uniquement lorsque la relation est autorisee ;",
        "l'administrateur peut bloquer un utilisateur et traiter les signalements ;",
        "les offres d'une entreprise refusee, bloquee ou non validee ne doivent pas etre visibles publiquement.",
    ])

    add_heading(doc, "10. Matching et suggestions")
    add_para(doc, "Le matching permet d'estimer la compatibilite entre une offre et un candidat. Il repose sur les informations du profil, les competences, la localisation, le niveau d'etude et les informations disponibles dans le CV lorsque celles-ci sont exploitables.")
    add_para(doc, "L'objectif n'est pas de remplacer le recruteur, mais de l'aider a prioriser les profils et d'aider le candidat a identifier les offres les plus pertinentes.")

    add_heading(doc, "11. Besoins non fonctionnels")
    for sub, items in [
        ("11.1 Securite", ["protection des mots de passe ;", "controle des acces selon les roles ;", "protection contre les actions non autorisees ;", "stockage prive des CV ;", "validation des donnees saisies ;", "limitation des tentatives sensibles ;", "moderation des messages et signalements."]),
        ("11.2 Ergonomie", ["interface claire pour chaque type d'utilisateur ;", "navigation simple entre les principales fonctionnalites ;", "messages d'erreur et de confirmation comprehensibles ;", "tableaux de bord lisibles ;", "formulaires organises et faciles a remplir."]),
        ("11.3 Performance et fiabilite", ["listes paginees lorsque les donnees sont nombreuses ;", "temps de reponse acceptable pour les operations courantes ;", "conservation correcte des donnees ;", "donnees de demonstration pour tester les parcours."]),
        ("11.4 Maintenabilite", ["organisation claire du code ;", "separation des responsabilites ;", "facilite d'evolution des fonctionnalites ;", "validation du comportement par des tests."]),
    ]:
        add_heading(doc, sub, 2)
        add_bullets(doc, items)

    add_heading(doc, "12. Contraintes du projet")
    add_bullets(doc, [
        "le projet doit etre une application web utilisable en local pour la demonstration ;",
        "la base de donnees doit stocker les comptes, profils, offres, candidatures, messages et signalements ;",
        "les documents sensibles, notamment les CV, ne doivent pas etre exposes publiquement ;",
        "la solution doit etre compatible avec le contexte de formation en developpement web et mobile ;",
        "les fonctionnalites doivent rester demonstrables pendant la soutenance.",
    ])

    add_heading(doc, "13. Criteres d'acceptation")
    add_bullets(doc, [
        "un visiteur peut consulter les offres publiques ;",
        "un candidat peut creer un compte, completer son profil, ajouter un CV et postuler ;",
        "une entreprise validee peut publier une offre et gerer les candidatures ;",
        "une entreprise non validee ne peut pas publier d'offre ;",
        "le systeme propose des suggestions ou scores de matching ;",
        "les CV ne sont accessibles qu'aux utilisateurs autorises ;",
        "la messagerie respecte les regles de controle ;",
        "l'administrateur peut gerer les utilisateurs, entreprises, referentiels et signalements ;",
        "les principaux parcours peuvent etre presentes sans erreur lors de la demonstration.",
    ])

    add_heading(doc, "14. Planning previsionnel")
    add_table(doc, ["Phase", "Contenu", "Etat"], [
        ("Phase 1", "Analyse du besoin et conception", "Terminee"),
        ("Phase 2", "Mise en place du projet et authentification", "Terminee"),
        ("Phase 3", "Espace candidat et gestion du CV", "Terminee"),
        ("Phase 4", "Espace entreprise et gestion des offres", "Terminee"),
        ("Phase 5", "Espace administrateur et moderation", "Terminee"),
        ("Phase 6", "Matching, suggestions et notifications", "Terminee"),
        ("Phase 7", "Tests, corrections et documentation", "En cours"),
        ("Phase 8", "Preparation de la soutenance et finalisation", "En cours"),
    ], [3.0, 9.5, 2.5])

    add_heading(doc, "15. Evolutions possibles")
    add_bullets(doc, [
        "amelioration du score de matching ;",
        "recherche avancee multicritere ;",
        "calendrier d'entretien ;",
        "historique detaille des actions administratives ;",
        "export PDF avance du CV ;",
        "application mobile ;",
        "statistiques plus detaillees pour les entreprises.",
    ])

    add_heading(doc, "16. Conclusion")
    add_para(doc, "JobConnect vise a fournir une plateforme de recrutement complete, securisee et adaptee aux besoins des candidats, des entreprises et de l'administrateur. Le projet met l'accent sur la qualite des parcours utilisateurs, la protection des donnees, la moderation et l'aide a la decision grace au matching.")
    add_para(doc, "Ce cahier des charges servira de reference pour verifier que la solution realisee repond bien aux besoins fonctionnels et non fonctionnels definis.")

    doc.save(OUT_DOCX)
    OUT_MD.write_text(build_markdown(), encoding="utf-8")
    build_pdf(build_markdown())
    print(OUT_DOCX)
    print(OUT_MD)
    print(OUT_PDF)


def build_pdf(markdown):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=12.5,
        spaceAfter=6,
        textColor=colors.HexColor("#1A202C"),
    )
    h1 = ParagraphStyle(
        "H1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        spaceBefore=12,
        spaceAfter=7,
        textColor=colors.HexColor("#1F4E79"),
    )
    h2 = ParagraphStyle(
        "H2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        spaceBefore=8,
        spaceAfter=5,
        textColor=colors.HexColor("#1F4E79"),
    )
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=26,
        leading=31,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1F4E79"),
        spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "Subtitle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=18,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#1A202C"),
        spaceAfter=6,
    )
    muted_center = ParagraphStyle(
        "MutedCenter",
        parent=body,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#5B6770"),
        fontSize=9,
        leading=12,
    )
    bullet_style = ParagraphStyle(
        "Bullet",
        parent=body,
        leftIndent=14,
        firstLineIndent=-8,
        spaceAfter=3,
    )

    def clean_inline(text):
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        while "**" in text:
            text = text.replace("**", "<b>", 1).replace("**", "</b>", 1)
        return text

    story = [
        Spacer(1, 2.7 * cm),
        Paragraph("JobConnect", title_style),
        Paragraph("Cahier des charges fonctionnel", subtitle_style),
        Paragraph("Plateforme web de recrutement pour candidats et entreprises", muted_center),
        Spacer(1, 0.8 * cm),
    ]
    meta = [
        ["Information", "Valeur"],
        ["Candidat", "Saad Berra"],
        ["Etablissement", "SupTechnology"],
        ["Filiere", "Developpement Web et Mobile"],
        ["Annee universitaire", "2025-2026"],
        ["Encadrant", "M. Ahmed Zellou"],
        ["Type de projet", "Application web de recrutement"],
        ["Objectif du document", "Definir le besoin, le perimetre, les acteurs, les fonctionnalites attendues, les contraintes et les criteres de validation."],
    ]
    story.append(make_pdf_table(meta, [4.2 * cm, 11.6 * cm], body))
    story.append(Spacer(1, 0.4 * cm))
    story.append(Paragraph("Version corrigee selon les remarques de l'encadrant : contenu recentre sur le cahier des charges, details techniques reduits.", muted_center))
    story.append(PageBreak())

    lines = markdown.splitlines()
    in_table = False
    table_lines = []
    skip_cover = True
    for raw in lines:
        line = raw.strip()
        if skip_cover:
            if line == "## 1. Introduction":
                skip_cover = False
            else:
                continue

        if line.startswith("|"):
            in_table = True
            table_lines.append(line)
            continue
        if in_table:
            add_markdown_table(story, table_lines, body)
            in_table = False
            table_lines = []

        if not line:
            continue
        if line.startswith("## "):
            story.append(Paragraph(clean_inline(line[3:]), h1))
        elif line.startswith("### "):
            story.append(Paragraph(clean_inline(line[4:]), h2))
        elif line.startswith("- "):
            story.append(Paragraph("• " + clean_inline(line[2:]), bullet_style))
        elif line.startswith("# "):
            continue
        else:
            story.append(Paragraph(clean_inline(line), body))
    if table_lines:
        add_markdown_table(story, table_lines, body)

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=A4,
        rightMargin=1.7 * cm,
        leftMargin=1.7 * cm,
        topMargin=1.7 * cm,
        bottomMargin=1.6 * cm,
        title="Cahier des charges JobConnect corrige",
        author="Saad Berra",
    )
    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)


def make_pdf_table(data, col_widths, body_style):
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Table, TableStyle

    wrapped = []
    for row in data:
        wrapped.append([Paragraph(str(cell), body_style) for cell in row])
    table = Table(wrapped, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    return table


def add_markdown_table(story, table_lines, body_style):
    from reportlab.lib.units import cm
    from reportlab.platypus import Spacer

    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(set(cell) <= {"-", ":", " "} for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    cols = len(rows[0])
    if cols == 2:
        widths = [4.0 * cm, 11.8 * cm]
    elif cols == 3:
        widths = [3.6 * cm, 8.7 * cm, 3.5 * cm]
    else:
        widths = [15.8 * cm / cols] * cols
    story.append(make_pdf_table(rows, widths, body_style))
    story.append(Spacer(1, 0.25 * cm))


def draw_footer(canvas, doc):
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#5B6770"))
    canvas.drawString(doc.leftMargin, 0.9 * cm, "JobConnect - Cahier des charges corrige")
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.9 * cm, f"Page {doc.page}")
    canvas.restoreState()


if __name__ == "__main__":
    build_docx()
